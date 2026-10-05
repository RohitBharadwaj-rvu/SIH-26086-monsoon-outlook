"""Rebuild the SIH26074 idea-submission deck on the team's original template (FAST&CURIOUS_SIH.pdf layout),
with the content updated to the final v3 system. All coordinates are in the original PDF's units (1440 x 810).

  python docs/deck074_v2/build_deck.py  -> docs/deck074_v2/SIH26074_FAST_CURIOUS.pptx
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

D = Path(__file__).resolve().parent
A = D.parents[1] / "docs" / "deck074_assets"

GREEN_FILL, GREEN_LINE, GREEN_LIGHT = "e9f6ed", "0e803d", "dff0df"
ORANGE_FILL, ORANGE = "fff2eb", "f26f20"
DARK_GREEN, DEEP_GREEN, BLUE, PURPLE = "054e3b", "0a6e4e", "006fbf", "8064a2"
TNR, ARIAL = "Times New Roman", "Arial"


def U(v):
    return Emu(int(round(v * 12192000 / 1440)))


def PT(sz):                       # original font size (PDF units) -> points
    return Pt(sz * 960 / 1440)


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = U(1440), U(810)
BLANK = prs.slide_layouts[6]


def shape(sl, kind, x, y, w, h, fill=None, line=None, lw=6, radius=None, dash=False):
    s = sl.shapes.add_shape(kind, U(x), U(y), U(w), U(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = PT(lw)
        if dash:
            s.line.dash_style = 4          # dash
    else:
        s.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def card(sl, x, y, w, h, color="green", radius=0.09, fill=None):
    f, l = {"green": (GREEN_FILL, GREEN_LINE), "orange": (ORANGE_FILL, ORANGE),
            "light": (GREEN_LIGHT, GREEN_LINE), "stage": (GREEN_FILL, DARK_GREEN)}[color]
    return shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill or f, l, 6, radius)


def text(sl, x, y, w, h, paras, font=TNR, size=24, color="000000", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bullet=False, spacing=1.0, space_after=0):
    """paras: list of paragraphs; each paragraph is a string or a list of (text, bold) runs."""
    tb = sl.shapes.add_textbox(U(x), U(y), U(w), U(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = [(para, False)] if isinstance(para, str) else para
        for t, b in runs:
            r = p.add_run(); r.text = t
            r.font.name = font; r.font.size = PT(size); r.font.bold = b; r.font.color.rgb = rgb(color)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(U(26)))); pPr.set("indent", str(-int(U(22))))
            bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•")
    return tb


def arrow(sl, x1, y1, x2, y2, color="000000", lw=3, head=True, dash=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, U(x1), U(y1), U(x2), U(y2))
    c.line.color.rgb = rgb(color); c.line.width = PT(lw)
    if dash:
        c.line.dash_style = 4
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def dot(sl, x, y, color, r=7):
    return shape(sl, MSO_SHAPE.OVAL, x - r, y - r, 2 * r, 2 * r, color)


def picture(sl, path, x, y, w, h, border=True, fit=True):
    if border:
        shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, "ffffff", "000000", 8)
    iw, ih = Image.open(path).size
    if fit:
        s = min(w / iw, h / ih); pw, ph = iw * s, ih * s
        sl.shapes.add_picture(str(path), U(x + (w - pw) / 2), U(y + (h - ph) / 2), U(pw), U(ph))
    else:
        sl.shapes.add_picture(str(path), U(x), U(y), U(w), U(h))


def frame(sl, n, title, badge=(43, 7), title_size=54, title_y=32):
    """Common template: SIH logo, team oval, title, blue footer with page number."""
    sl.shapes.add_picture(str(D / "logo.png"), U(1155), U(0), U(266), U(126))
    o = shape(sl, MSO_SHAPE.OVAL, badge[0], badge[1], 151, 98, "ffffff", PURPLE, 6)
    tf = o.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, t in enumerate(("FAST&", "CURIOUS")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t; r.font.name = "Calibri"; r.font.size = PT(27); r.font.color.rgb = rgb("000000")
    text(sl, 220, title_y - 4, 930, 80, [[(title, True)]], size=title_size, align=PP_ALIGN.CENTER)
    shape(sl, MSO_SHAPE.RECTANGLE, 0, 751, 1440, 59, BLUE)
    text(sl, 520, 761, 400, 30, ["@SIH Idea submission"], font=ARIAL, size=18, color="ffffff", align=PP_ALIGN.CENTER)
    text(sl, 1320, 761, 50, 30, [[(str(n), True)]], font=ARIAL, size=18, color="ffffff", align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- slide 1: title page (unchanged content)
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(str(D / "p1_graphic.png"), U(668), U(101), U(549), U(609))
s.shapes.add_picture(str(D / "logo.png"), U(1155), U(0), U(266), U(126))
text(s, 100, 18, 1100, 90, [[("SMART INDIA HACKATHON 2026", True)]], font="Garamond", size=60, color="1f497d",
     align=PP_ALIGN.CENTER)
items = ["Problem Statement ID – SIH26074",
         "Problem Statement Title - Downscaling of Weather Forecast from Block level to Panchayat level",
         "Theme- Agriculture,FoodTech & Rural Development", "PS Category- Software",
         "Team ID - 173900", "Team Name - FAST&CURIOUS"]
text(s, 64, 185, 1350, 570, [[(t, True)] for t in items], font=ARIAL, size=36, bullet=True, spacing=1.0, space_after=26)

# ---------------------------------------------------------------- slide 2: proposed solution
s = prs.slides.add_slide(BLANK)
frame(s, 2, "Spatiotemporal Diffusion Downscaler", badge=(56, 16), title_size=36, title_y=50)
text(s, 36, 150, 1390, 80, [[("Problem Statement: A coarse 27 km (0.25°) GFS forecast gives one value for dozens of Gram "
                               "Panchayats, missing local showers and terrain effects.", True)]],
     size=30)
text(s, 37, 243, 400, 36, [[("Proposed Solution:", True)]], size=26)
bul = [
    [("Spatiotemporal Transformer Downscaling", True), (": attention across space, time and the 6 variables, with a "
     "Mixture-of-Experts layer, turns the coarse forecast into a 0.05° (~5.5 km) field.", False)],
    [("Cross-Fitted Residual Diffusion", True), (": generates an ensemble of fine-scale detail on top of the transformer; "
     "each training season comes from a model that never saw it (rain CRPS −19 %).", False)],
    [("Topography Coupling", True), (": Copernicus GLO-30 elevation layers and a land mask capture ridge and valley effects.", False)],
    [("Calibrated Uncertainty", True), (": spread calibration and a rain tail cap; 5–95 % ranges cover 88 % of outcomes "
     "on the unseen 2023 monsoon.", False)],
    [("Real Forecasts, Real Skill", True), (": trained on 1,096 genuine GFS forecasts; on 2023 it cuts wet-day rain error by "
     "20 % and Tmax error by 24 % versus raw GFS.", False)],
]
text(s, 60, 292, 750, 450, bul, size=24, bullet=True, spacing=0.98, space_after=5)
# flowchart (right)
card(s, 833, 208, 555, 79, "green", 0.18)
text(s, 840, 216, 540, 64, [[("INPUTS", True)], [("Real GFS 0.25° (7 days × 6 vars), ERA5 history, Terrain", True)]],
     font=ARIAL, size=17, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
card(s, 834, 318, 555, 79, "green", 0.18)
text(s, 840, 326, 540, 64, [[("TRANSFORMER BACKBONE", True)], [("Spatiotemporal attention + Mixture-of-Experts", True)]],
     font=ARIAL, size=17, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
card(s, 833, 427, 555, 79, "green", 0.18)
text(s, 840, 435, 540, 64, [[("MACRO BASELINE B(X)", True)], [("Deterministic forecast, 3 seeds averaged (FAST)", True)]],
     font=ARIAL, size=17, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1001, 527, 387, 99, ORANGE, "c2400c", 6, 0.18)
text(s, 1008, 534, 373, 85, [[("RESIDUAL DIFFUSION", True)], [("Cross-fitted, DPM-Solver++ 24–32 steps", True)]],
     font=ARIAL, size=17, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 831, 545, 138, 69, ORANGE_FILL, ORANGE, 4, 0.18, dash=True)
text(s, 835, 549, 130, 61, [[("Noise:", True)], [("z_T ~ N(0, I)", True)]], font=ARIAL, size=17, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 833, 648, 555, 91, DEEP_GREEN, DARK_GREEN, 6, 0.18)
text(s, 840, 652, 540, 84, [[("PROBABILISTIC FORECAST: Y = B(X) + R", True)],
                            [("7-Day × 6-Variable (P, Tmax, Tmin, RH, U, V) → 234 Gram Panchayats", True)]],
     font=ARIAL, size=17, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for y1, y2 in ((287, 316), (397, 425), (506, 525), (626, 646)):
    arrow(s, 1111, y1, 1111, y2)
arrow(s, 969, 580, 999, 580)

# ---------------------------------------------------------------- slide 3: technical approach
s = prs.slides.add_slide(BLANK)
frame(s, 3, "TECHNICAL APPROACH", badge=(43, 7))
text(s, 21, 148, 290, 30, [[("TECH STACK", True)]], size=22)
stack = ["PyTorch", "Python, NumPy", "Xarray, Zarr, NetCDF", "ecCodes (GRIB)", "FastAPI, Uvicorn", "Pydantic",
         "Shapely", "Leaflet.js", "HTML5, CSS, JS", "Service Worker (PWA)", "SQLite", "Kaggle T4 GPUs"]
text(s, 30, 176, 285, 400, stack, size=22, bullet=True, spacing=0.92)
text(s, 21, 585, 290, 60, [[("PERFORMANCE &", True)], [("TESTING", True)]], size=22, spacing=0.92)
text(s, 30, 640, 285, 110, ["0.17 s FAST forecast (T4)", "8.3 s for 16 members", "< 1 GB GPU memory",
                            "234 GPs × 7 days × 6 vars"], size=22, bullet=True, spacing=0.92)
arrow(s, 318, 164, 318, 675, head=False)
# stage 1
card(s, 355, 127, 280, 565, "light", 0.06)
text(s, 355, 140, 280, 26, [[("STAGE 1 - INPUTS", True)]], size=20, align=PP_ALIGN.CENTER)
for (yy, h, icon, lines) in ((180, 156, "icon_clock.png", ["Historical Weather", "(ERA5, last H = 3 days,", "actual obs up to t0)"]),
                             (363, 140, "icon_sat.png", ["Future Coarse NWP", "(real GFS 0.25°,", "D+1 … D+7)"]),
                             (534, 140, "icon_terrain.png", ["Terrain & Context", "(Copernicus GLO-30 DEM", "layers + land mask)"])):
    card(s, 382, yy, 234, h, "stage", 0.1)
    s.shapes.add_picture(str(D / icon), U(472), U(yy + 8), U(56), U(48))
    text(s, 388, yy + 62, 222, h - 66, [[(l, True)] for l in lines], size=18, align=PP_ALIGN.CENTER, spacing=0.95)
# stage 2
card(s, 672, 126, 474, 160, "stage", 0.08)
text(s, 680, 136, 458, 145, [[("STAGE 2 - SPATIOTEMPORAL BACKBONE", True)], [("Transformer + Cross-Attention", True)],
                             [("(History <-> Future Conditioning)", False)],
                             [("Mixture-of-Experts: ", True), ("8 experts, top-2 routing", False)],
                             [("Variables: ", True), ("P, T_max, T_min, RH, U, V", False)],
                             [("Deterministic Macro Baseline B(X)", True)]],
     size=18, align=PP_ALIGN.CENTER, spacing=0.92)
# stage 3
card(s, 679, 296, 476, 186, "orange", 0.08)
text(s, 687, 306, 460, 170, [[("STAGE 3 - CROSS-FITTED RESIDUAL DIFFUSION", True)],
                             [("Gaussian Initialization: ", True), ("z_T ~ N(0, I)", False)],
                             [("Cross-Fitting: ", True), ("each season's residuals from a model that never saw it", False)],
                             [("Reverse Denoising: ", True), ("DPM-Solver++, 24–32 steps", False)],
                             [("Ensemble: ", True), ("8–16 members, R(1) … R(E)", False)]],
     size=18, align=PP_ALIGN.CENTER, spacing=0.95, space_after=3)
# stage 4
card(s, 681, 510, 475, 214, "stage", 0.07)
text(s, 689, 518, 459, 200, [[("STAGE 4 - CALIBRATION & POST-PROCESSING", True)],
                             [("Y = ", True), ("B(X) + R (summation per ensemble member)", False)],
                             [("Calibration: ", True), ("spread per lead × variable, rain quantile mapping, rain tail cap", False)],
                             [("Physical Inversion: ", True), ("expm1(precip), unstandardize physical vars", False)],
                             [("Zonal GIS Aggregation: ", True), ("0.05° grid → 234 Gram Panchayats of Mandya (area-weighted)", False)]],
     size=18, align=PP_ALIGN.CENTER, spacing=0.95, space_after=3)
# phone (Agro-Weather PWA)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1199, 127, 167, 331, "e8e8e8", "1e1e1e", 9, 0.12)
text(s, 1207, 140, 151, 312, [[("Agro-Weather PWA", True)], "", "7-Day × 6-Variable Panchayat Forecast", "",
                              "5–95 % Ranges & Exceedance Probabilities", "", "4 Modes: FAST to 16-member ENSEMBLE", "",
                              "Proof plots: GFS vs v3 vs observed"],
     size=15, align=PP_ALIGN.CENTER, spacing=0.9)
# validation
card(s, 1182, 495, 216, 227, "orange", 0.08)
text(s, 1188, 502, 204, 215, [[("PROBABILISTIC EVALUATION & VALIDATION", True)], [("Held-out Test", True)],
                              [("(2023 monsoon, 122 forecasts; trained 2015–22)", False)],
                              [("Metrics: ", True), ("CRPS, spread-skill ratio, coverage, CSS", False)]],
     size=17, align=PP_ALIGN.CENTER, spacing=0.92, space_after=3)
# connectors
arrow(s, 651, 206, 651, 617, GREEN_LINE, 3, head=False)
for yy, col in ((206, DEEP_GREEN), (617, DEEP_GREEN)):
    arrow(s, 651, yy, 670, yy, GREEN_LINE, 3); dot(s, 651, yy, col)
arrow(s, 619, 431, 677, 431, ORANGE, 3); dot(s, 619, 431, DEEP_GREEN)
arrow(s, 908, 287, 908, 294, GREEN_LINE, 3)
arrow(s, 908, 483, 908, 508, GREEN_LINE, 3)
arrow(s, 1148, 229, 1197, 229, GREEN_LINE, 3, dash=True); dot(s, 1146, 229, DEEP_GREEN)
arrow(s, 1157, 418, 1170, 418, ORANGE, 3, head=False, dash=True); dot(s, 1155, 418, ORANGE)
arrow(s, 1170, 418, 1170, 493, ORANGE, 3, dash=True)
arrow(s, 1290, 493, 1290, 460, "000000", 3, dash=True)
text(s, 360, 697, 290, 50, [[("Note: Real GFS forecasts only - no future leakage", True)]], size=18, color="c00000",
     align=PP_ALIGN.CENTER)
text(s, 681, 727, 520, 24, [[("Note: Every setting chosen on 2022 or out-of-fold, never on 2023", True)]], size=18,
     color="c00000", align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- slide 4: feasibility and viability
s = prs.slides.add_slide(BLANK)
frame(s, 4, "FEASIBILITY AND VIABILITY", badge=(37, 10))
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 202, 108, 308, 74, DARK_GREEN, DARK_GREEN, 6, 0.5)
text(s, 202, 108, 308, 74, [[("FEASIBILITY", True)]], size=32, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 868, 102, 308, 74, ORANGE, ORANGE, 6, 0.5)
text(s, 868, 102, 308, 74, [[("VIABILITY", True)]], size=32, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 715, 175, 715, 686, head=False)
feas = [(209, 80, "green", "Runs in 0.17 s (FAST) to 8.3 s (16-member ensemble) on one T4 GPU, using < 1 GB of memory."),
        (298, 107, "orange", "Copernicus GLO-30 elevation layers and a land mask condition the model, capturing terrain effects at panchayat scale."),
        (415, 107, "light", "Uses free, operational GFS forecasts as the input layer, avoiding the need to build a new numerical weather prediction system.")]
for yy, h, c, t in feas:
    card(s, 72, yy, 607, h, c, 0.12)
    text(s, 92, yy + 6, 570, h - 12, [[(t, True)]], size=24, bullet=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
picture(s, A / "app_timeline.png", 72, 548, 607, 148)
viab = [(209, 80, "orange", "Generates fine-resolution 7-day, 6-variable forecasts with calibrated uncertainty ranges."),
        (298, 80, "green", "Converts GFS into Panchayat-level forecasts, maps and exceedance probabilities."),
        (388, 80, "orange", "Supports localized agro-weather advisories, Panchayat maps and spatial risk visualization."),
        (480, 80, "green", "Input-agnostic: NCMRWF or IMD forecasts can replace GFS after retraining on their archives."),
        (572, 135, "orange", "Scales to any region with the same global data (GFS, CHIRPS, ERA5-Land, GLO-30) by retraining; the current model and demo cover all 234 Gram Panchayats of Mandya.")]
for yy, h, c, t in viab:
    card(s, 752, yy, 607, h, c, 0.12)
    text(s, 772, yy + 6, 570, h - 12, [[(t, True)]], size=24, bullet=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)

# ---------------------------------------------------------------- slide 5: impact and benefits
s = prs.slides.add_slide(BLANK)
frame(s, 5, "IMPACT AND BENEFITS", badge=(39, 0))
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 218, 98, 220, 51, DARK_GREEN, DARK_GREEN, 6, 0.5)
text(s, 218, 98, 220, 51, [[("IMPACT", True)]], size=24, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 220, 451, 218, 52, ORANGE, ORANGE, 6, 0.5)
text(s, 220, 451, 218, 52, [[("BENEFITS", True)]], size=24, color="ffffff", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
imp = [(159, 99, "green", [("Panchayat-level forecasts for all 234 Gram Panchayats of Mandya, validated on the unseen 2023 monsoon.", True)]),
       (268, 80, "orange", [("Hyper-local weather information (0.25° → 0.05°)", True)]),
       (358, 80, "green", [("Beats raw GFS at Panchayat level: wet-day rain error −20 %, Tmax error −24 %.", True)]),
       (516, 99, "orange", [("6-Variable, 7-Day Forecasts: Joint precipitation, temperature, humidity and wind information "
                             "provides richer agricultural context than rainfall alone.", True)]),
       (625, 99, "green", [("Calibrated Uncertainty: 5–95 % ranges cover 88 % of outcomes, and heavy-rain maps double the "
                            "detection of ≥ 64.5 mm days.", True)])]
for yy, h, c, runs in imp:
    card(s, 39, yy, 607, h, c, 0.12)
    text(s, 59, yy + 6, 575, h - 12, [runs], size=22, bullet=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
picture(s, D / "proof_slide5.png", 738, 149, 654, 291)
picture(s, D / "app_crop.png", 738, 448, 654, 285)

# ---------------------------------------------------------------- slide 6: research and references
s = prs.slides.add_slide(BLANK)
frame(s, 6, "RESEARCH AND REFERENCES", badge=(37, 10))
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 81, 143, 584, 62, DARK_GREEN, DARK_GREEN, 6, 0.5)
text(s, 81, 143, 584, 62, [[("DATASETS & PROJECT BASES", True)]], size=28, color="ffffff", align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 761, 148, 598, 56, ORANGE, ORANGE, 6, 0.5)
text(s, 761, 148, 598, 56, [[("KEY METHODOLOGICAL REFERENCES", True)]], size=28, color="ffffff", align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 724, 174, 724, 685, head=False)
data = ["NOAA GFS 0.25° operational forecasts (AWS Open Data, NCAR RDA ds084.1); model input.",
        "CHIRPS v2.0 0.05° daily precipitation (Funk et al., 2015); high-res rain target.",
        "Copernicus DEM GLO-30 30m DEM; terrain conditioning (elevation layers, land mask).",
        "ERA5 history & ERA5-Land 0.1° (Tmax, Tmin, RH, U, V) targets; Hersbach et al. 2020.",
        "Panchayat Boundaries & LGD Codes for Mandya, Karnataka (234 Gram Panchayats)."]
refs = ["Mardani et al. (2025) CorrDiff: residual corrective diffusion for km-scale downscaling. Commun. Earth Environ.",
        "Ho, Jain & Abbeel (2020) DDPM: basis for the residual diffusion component. NeurIPS 2020.",
        "Lu et al. (2022) DPM-Solver++: fast sampling of diffusion models. arXiv:2211.01095.",
        "Fedus, Zoph & Shazeer (2022) Switch Transformers: sparse Mixture-of-Experts. JMLR 23.",
        "Vandal et al. (2017) DeepSD: High-res climate projections via image super-resolution. KDD 2017."]
ys = (228, 323, 420, 514, 609)
for i, (yy, t) in enumerate(zip(ys, data)):
    card(s, 81, yy, 607, 80, "green" if i % 2 == 0 else "orange", 0.12)
    text(s, 101, yy + 6, 575, 68, [[(t, True)]], size=22, bullet=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
for i, (yy, t) in enumerate(zip(ys, refs)):
    card(s, 761, yy, 607, 80, "orange" if i % 2 == 0 else "green", 0.12)
    text(s, 781, yy + 6, 575, 68, [[(t, True)]], size=22, bullet=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
text(s, 81, 703, 1287, 30, [[("Code, model weights & reports: ", True),
                             ("github.com/rohzhegde26/SIH-26074-Fast-and-Curious", False)]], size=20,
     align=PP_ALIGN.CENTER, color="0e4f8a")

out = D / "SIH26074_FAST_CURIOUS.pptx"
prs.save(out)
print("saved", out)
